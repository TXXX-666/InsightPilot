from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.shared import Pt
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer


def markdown_to_docx(markdown: str, output: Path) -> Path:
    document = Document()
    styles = document.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"].font.size = Pt(10.5)
    for raw in markdown.splitlines():
        line = raw.strip()
        if not line:
            document.add_paragraph()
        elif line.startswith("### "):
            document.add_heading(line[4:], level=3)
        elif line.startswith("## "):
            document.add_heading(line[3:], level=2)
        elif line.startswith("# "):
            document.add_heading(line[2:], level=1)
        elif line.startswith(("- ", "* ")):
            document.add_paragraph(line[2:], style="List Bullet")
        elif re.match(r"^\d+\.\s", line):
            document.add_paragraph(re.sub(r"^\d+\.\s", "", line), style="List Number")
        else:
            document.add_paragraph(line)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    return output


def markdown_to_pdf(markdown: str, output: Path) -> Path:
    output.parent.mkdir(parents=True, exist_ok=True)
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    styles = getSampleStyleSheet()
    body = ParagraphStyle("ChineseBody", parent=styles["BodyText"], fontName="STSong-Light", fontSize=10, leading=16, alignment=TA_LEFT, wordWrap="CJK")
    heading = ParagraphStyle("ChineseHeading", parent=body, fontSize=16, leading=22, spaceBefore=8, spaceAfter=6)
    story = []
    for raw in markdown.splitlines():
        line = raw.strip()
        if not line:
            story.append(Spacer(1, 4 * mm))
            continue
        is_heading = line.startswith("#")
        clean = re.sub(r"^#{1,6}\s*", "", line)
        clean = clean.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        story.append(Paragraph(clean, heading if is_heading else body))
    doc = SimpleDocTemplate(str(output), pagesize=A4, rightMargin=18 * mm, leftMargin=18 * mm, topMargin=18 * mm, bottomMargin=18 * mm, title="InsightPilot Research Report")
    doc.build(story)
    return output

